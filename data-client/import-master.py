#!/usr/bin/env python3
"""Import master data (salesman, customer sekolah, item) dari xlsx data-client ke DB ERP.

Semua baris diberi legacy_code='data-client' -> rollback: lihat ROLLBACK_SQL di bawah.
Pakai: python3 import-master.py [--apply]   (tanpa --apply = dry-run / rollback di akhir)
"""
import collections, json, os, re, subprocess, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TAG = 'data-client'
APPLY = '--apply' in sys.argv

KEC_CODE = {
    'KEDUNGKANDANG': '3573010', 'SUKUN': '3573020', 'KLOJEN': '3573030', 'BLIMBING': '3573040',
    'LOWOKWARU': '3573050', 'PAKIS': '3507250', 'PAKISAJI': '3507220', 'WAGIR': '3507210',
    'SINGOSARI': '3507280', 'LAWANG': '3507270', 'KARANGPLOSO': '3507290', 'BUMIAJI': '3579030',
    'TUMPANG': '3507240', 'PONCOKUSUMO': '3507100', 'BULULAWANG': '3507130', 'WAJAK': '3507110',
}
KEC_ALIAS = {'SAKUN': 'SUKUN'}

SCHOOL_ALIAS = {
    'TL ABA 5': 'TK ABA 5', 'ABA 5': 'TK ABA 5', 'ABA 24': 'TK ABA 24',
    'TK A YANI': 'TK AHMAD YANI', 'WONDERFULL KIDS': 'TK WONDERFULL KIDS',
    'TK CENDRAWASI': 'TK CENDRAWASIH', 'TK MASLIMAT NU 34': 'TK MUSLIMAT NU 34',
    'PP PERMATA KASIH': 'TK PERMATA KASIH', 'PP. ROSALIA': 'POS PAUD ROSALIA',
    'RA NA NAYARA': 'RA AN NAYARA', 'TK NAYARA': 'RA AN NAYARA',
    'TK MUSLIMAT NU 28': 'TK MUSLIMAT NU 28 HASYIM ASYARI',
    'TK MUSLIMAT NU 28 HASYIM': 'TK MUSLIMAT NU 28 HASYIM ASYARI',
    "TK HASYIM ASY'ARI": 'TK MUSLIMAT NU 28 HASYIM ASYARI',
    'TK AR ROHMAN 9': 'TK AR ROHMAN 9',
}
SKIP_AMBIGUOUS = {'TK DWP', 'TK TRISULA'}
SALESMEN = {  # key sheet -> (nama, phone)
    'VIKRI': ('M. VIKRI', '087756166777'), 'IVAN': ('IVAN', None), 'BAMBANG': ('BAMBANG', None),
}
FLAGS = collections.defaultdict(list)


def sq(s):
    return 'NULL' if s is None else "'" + str(s).replace("'", "''") + "'"


def norm_school(raw):
    n = re.sub(r'\s+', ' ', str(raw).replace('\xa0', ' ')).strip().upper()
    n = re.sub(r'\s*\(\d{4}\)$', '', n)
    return SCHOOL_ALIAS.get(n, n)


def norm_kec(raw):
    if not raw:
        return None
    k = re.sub(r'\s+', ' ', str(raw)).strip().upper()
    return KEC_ALIAS.get(k, k)


def norm_phone(raw):
    if raw is None:
        return None
    d = re.sub(r'\D', '', str(int(raw)) if isinstance(raw, float) else str(raw))
    if not d:
        return None
    if d.startswith('62'):
        d = '0' + d[2:]
    elif not d.startswith('0'):
        d = '0' + d
    return d


def rows(wb, sheet, maxc=16):
    return [list(r[:maxc]) for r in wb[sheet].iter_rows(values_only=True)]


custs = collections.OrderedDict()


def add_cust(raw, kec=None, contact=None, phone=None, salesman=None, src=''):
    name = norm_school(raw)
    if not name or name in SKIP_AMBIGUOUS:
        if name in SKIP_AMBIGUOUS:
            FLAGS['sekolah_dilewati_ambigu'].append(f'{name} ({src})')
        return
    c = custs.setdefault(name, {'kec': None, 'contacts': [], 'salesman': None, 'src': set()})
    c['src'].add(src)
    k = norm_kec(kec)
    if k and not c['kec']:
        c['kec'] = k
    if salesman and not c['salesman']:
        c['salesman'] = salesman
    ph = norm_phone(phone)
    cn = re.sub(r'\s+', ' ', str(contact)).strip().upper() if contact else None
    if cn or ph:
        same = [i for i, x in enumerate(c['contacts']) if x[0] == cn]
        if same and (ph is None or c['contacts'][same[0]][1] in (None, ph)):
            if ph:
                c['contacts'][same[0]] = (cn, ph)
        elif not any(x[0] == cn and x[1] == ph for x in c['contacts']):
            c['contacts'].append((cn or 'KONTAK', ph))


# --- marketing (prospek per salesman) ---
wb = openpyxl.load_workbook(os.path.join(HERE, 'LAPORAN MARKETING.xlsx'), read_only=True, data_only=True)
for sheet, sm in (('VIQRI', 'VIKRI'), ('IVAN', 'IVAN')):
    for r in rows(wb, sheet)[1:]:
        if r[3]:
            add_cust(r[3], r[4], r[5], r[6], sm, f'marketing/{sheet}')

# --- pembayaran ---
wb = openpyxl.load_workbook(os.path.join(HERE, 'LAPORAN PEMBAYARAN.xlsx'), read_only=True, data_only=True)
for sheet, sm in (('VIKRI', 'VIKRI'), ('IVAN', 'IVAN'), ('BAMBANG', 'BAMBANG')):
    for r in rows(wb, sheet)[2:]:
        if r[1]:
            add_cust(r[1], r[2], r[3], r[4], sm, f'pembayaran/{sheet}')

# --- PO / SJ / Invoice header ---
wb = openpyxl.load_workbook(os.path.join(HERE, 'KELENGKAPAN PEMESANAN.xlsx'), read_only=True, data_only=True)
inv_items, po_items = {}, []
for sheet in ('PO', 'SJ', 'INVOICE'):
    for r in rows(wb, sheet, 13):
        t = [x for x in r if x not in (None, '')]
        if not t:
            continue
        if t[0] == 'NAMA SEKOLAH':
            add_cust(str(t[1]).lstrip(': '), None, None, None, None, f'transaksi/{sheet}')
        if isinstance(t[0], str) and re.fullmatch(r'\d+\.', t[0].strip()) and len(t) >= 3:
            nm = str(t[1]).strip()
            if sheet == 'INVOICE' and len(t) >= 4 and isinstance(t[3], (int, float)):
                inv_items[nm] = t[3]
            else:
                po_items.append(nm)

# --- buku, map raport, tagihan ---
wb = openpyxl.load_workbook(os.path.join(HERE, 'PEMESANAN BUKU.xlsx'), read_only=True, data_only=True)
book_titles = []
for r in rows(wb, 'Sheet1')[1:]:
    if r[4]:
        add_cust(r[4], r[5], None, None, None, 'pemesanan-buku')
    if r[6] and str(r[6]).strip().upper() in {
        'ENSIKLOPEDIA SI KECIL ISI 12 BUKU KB', 'SEJARAH 25 NABI DAN RASUL', 'PEMBELAJARAN UNTUK FASE PONDASI',
        'BELAJAR DAN BARMAIN BERBASIS BUKU', 'NILAI AGAMA DAN BUDI PEKERTI', 'JATI DIRI',
        'PROJEK PENGUATAN PROFIL PELAJAR PANCASILA', 'DINO KIDS KB', 'BUDI SI PENGANTAR PAKET'}:
        book_titles.append(re.sub(r'\s+', ' ', str(r[6])).strip().upper())
wb = openpyxl.load_workbook(os.path.join(HERE, 'PEMESANAN MAP RAPORT.xlsx'), read_only=True, data_only=True)
map_items = {}
for r in rows(wb, 'Sheet1')[1:]:
    if r[4]:
        add_cust(r[4], None, None, None, None, 'map-raport')
    if r[6] and r[9]:
        map_items[re.sub(r'\s+', ' ', str(r[6])).strip().upper()] = r[9]
wb = openpyxl.load_workbook(os.path.join(HERE, 'LAPORAN KEUANGAN.xlsx'), read_only=True, data_only=True)
for r in rows(wb, 'TAGIHAN MARKETING', 8)[1:]:
    if r[4] and r[5] and str(r[0]).strip() not in ('TOTAL',):
        add_cust(r[4], None, None, None, None, 'tagihan-marketing')

# --- item ---
ITEM_ALIAS = {'ENSIKLOPEDIA SI KECIL ISI 12 BUKU': 'ENSIKLOPEDIA SI KECIL ISI 12 BUKU KB',
              'SEJARAH 25 NABI DAN RASAUL': 'SEJARAH 25 NABI DAN RASUL', 'TANTANGAN JAGARAGA': 'TANTANGAN JAGA RAGA',
              'BAHASA JAWA A': 'BUKU BAHASA JAWA A', 'BAHSA JAWA B': 'BUKU BAHASA JAWA B',
              'BAHASA JAWA B': 'BUKU BAHASA JAWA B', 'BUKU BAHASA JAWA': 'BUKU BAHASA JAWA',
              'TISSUE': 'TISSUE', 'BUKU HIJAIYAH AB': 'BUKU HIJAIYAH AB',
              'KERTAS BUFALLO PINK': 'KERTAS BUFALLO PINK'}


def norm_item(raw):
    n = re.sub(r'\s+', ' ', str(raw)).strip().upper()
    n = re.sub(r'\bBHS\b', 'BAHASA', n).replace('BAHSA', 'BAHASA')
    return ITEM_ALIAS.get(n, n)


BUY = {  # nama kanonik -> (harga beli terakhir, harga jual dari sheet pembelian bila invoice tak ada)
    'YARIS CHR 30CM RED': (115000, None), 'YARIS CHR 30CM BLUE': (115000, None),
    'YARIS CHR 30CM YLW': (115000, None), 'YARIS CHR 30CM GREEN': (115000, None),
    'YARIS TABLE RECTANGLE RED': (699000, None), 'YARIS TABLE RECTANGLE BLUE': (699000, None),
    'TINTA EPSON 003': (86000, None), 'KERTAS HVS A4': (37600, None), 'KERTAS HVS F4': (45500, None),
    'KERTAS BUFALLO PUTIH': (29000, None), 'KERTAS BUFALLO HIJAU PUPUS': (36000, None),
    'KERTAS BUFALLO ORANGE': (36000, None), 'KERTAS BUFALLO PINK': (24500, 33810),
    'KERTAS BUFALLO KUNING': (24500, 33810), 'KERTAS LIPAT 10X10': (1700, None),
    'KERTAS LIPAT 25X25': (7500, None), 'KERTAS ASTURO': (1600, None), 'CRAYON JOYKO': (17500, None),
    'MAP KANCING MERAH': (2575, None), 'SNOWMAN SPIDOL': (12700, None), 'TISSUE': (6300, None),
    'MATERAI': (10000, None), 'SABUN CUCI TANGAN': (26500, None), 'MINYAK KAYU PUTIH': (41603, None),
    'MINYAK TELON': (15965, None),
}
items = collections.OrderedDict()  # nama -> {sell, buy}
for nm, price in inv_items.items():
    i = items.setdefault(norm_item(nm), {'sell': None, 'buy': None})
    if price:
        if i['sell'] and i['sell'] != price:
            FLAGS['harga_jual_berbeda_antar_invoice'].append(f'{norm_item(nm)}: {i["sell"]} -> {price} (dipakai terakhir)')
        i['sell'] = price
for nm in po_items + book_titles + ['KERTAS BUFALLO KUNING']:
    items.setdefault(norm_item(nm), {'sell': None, 'buy': None})
for nm, price in map_items.items():
    items.setdefault(nm, {'sell': None, 'buy': None})['sell'] = price
for nm, (buy, sell) in BUY.items():
    i = items.setdefault(nm, {'sell': None, 'buy': None})
    i['buy'] = buy
    if not i['sell'] and sell:
        i['sell'] = sell

CAT_ATK = ('TINTA', 'KERTAS', 'CRAYON', 'MAP ', 'SPIDOL', 'MATERAI')
CAT_CON = ('TISSUE', 'SABUN', 'MINYAK')


def category(nm):
    if nm.startswith('YARIS'):
        return 'FURN'
    if any(nm.startswith(k) or k in nm for k in CAT_ATK):
        return 'ATK'
    if any(nm.startswith(k) for k in CAT_CON):
        return 'CON'
    return 'BUKU'


def unit(nm):
    if nm.startswith('KERTAS HVS'):
        return 'RIM'
    if nm.startswith('KERTAS') and 'HVS' not in nm:
        return 'LBR'
    return 'PCS'


def psql(sql, tuples=False):
    cmd = ['docker', 'exec', '-i', 'sentient-postgres-core', 'sh', '-c',
           'psql -U "${POSTGRES_USER:-postgres}" -d sentient_factory -v ON_ERROR_STOP=1 ' + ('-Atq' if tuples else '')]
    p = subprocess.run(cmd, input=sql, capture_output=True, text=True)
    if p.returncode:
        sys.exit('PSQL ERROR:\n' + p.stderr + p.stdout)
    return p.stdout


last = psql("select coalesce(max(substring(code from 6)::int),0) from md_partners where code ~ '^CUST-[0-9]+$' union all "
            "select coalesce(max(substring(code from 5)::int),0) from md_partners where code ~ '^SLS-[0-9]+$'", True).split()
cust_n, sls_n = int(last[0]), int(last[1])

sql = ['begin;']
sql.append(f"insert into md_item_categories(code,name,legacy_code,updated_at) values "
           f"('BUKU','Buku & Bahan Ajar','{TAG}',now()),('FURN','Furnitur Sekolah','{TAG}',now()) "
           f"on conflict do nothing;")
sm_code = {}
for key, (nm, ph) in SALESMEN.items():
    sls_n += 1
    sm_code[key] = f'SLS-{sls_n:04d}'
    sql.append(f"insert into md_partners(code,name,partner_type_id,legacy_code,updated_at,metadata) "
               f"select {sq(sm_code[key])},{sq(nm)},(select id from md_partner_types where code='SLS'),'{TAG}',now(),"
               f"{sq(json.dumps({'phone': ph}))}::jsonb where not exists (select 1 from md_partners where name={sq(nm)} and legacy_code='{TAG}');")
for name, c in custs.items():
    cust_n += 1
    code = f'CUST-{cust_n:04d}'
    smk = c['salesman']
    sm_sub = f"(select id from md_partners where legacy_code='{TAG}' and name={sq(SALESMEN[smk][0])})" if smk else 'NULL'
    sql.append(f"insert into md_partners(code,name,partner_type_id,salesman_id,legacy_code,updated_at) "
               f"select {sq(code)},{sq(name)},(select id from md_partner_types where code='CUST-SCHOOL'),{sm_sub},'{TAG}',now() "
               f"where not exists (select 1 from md_partners where name={sq(name)} and legacy_code='{TAG}');")
    pid = f"(select id from md_partners where name={sq(name)} and legacy_code='{TAG}')"
    if c['kec'] in KEC_CODE:
        sql.append(f"insert into md_partner_addresses(partner_id,type,is_default,address_line1,country_id,province_id,city_id,area_id,updated_at) "
                   f"select {pid},'OFFICE',true,{sq('Kec. ' + c['kec'].title())},p.country_id,c.province_id,c.id,a.id,now() "
                   f"from md_areas a join md_cities c on c.id=a.city_id join md_provinces p on p.id=c.province_id where a.code='{KEC_CODE[c['kec']]}' "
                   f"and not exists (select 1 from md_partner_addresses x where x.partner_id={pid});")
    elif c['kec']:
        FLAGS['kecamatan_tidak_dikenal'].append(f'{name}: {c["kec"]}')
    else:
        FLAGS['sekolah_tanpa_kecamatan'].append(name)
    if not any(x[1] for x in c['contacts']):
        FLAGS['sekolah_tanpa_telepon'].append(name)
    for idx, (cn, ph) in enumerate(c['contacts']):
        sql.append(f"insert into md_partner_contacts(partner_id,name,phone,is_default,updated_at) select {pid},{sq(cn)},{sq(ph)},{str(idx == 0).lower()},now() "
                   f"where not exists (select 1 from md_partner_contacts x where x.partner_id={pid} and x.name={sq(cn)} and coalesce(x.phone,'')=coalesce({sq(ph)},''));")
    if not c['salesman']:
        FLAGS['sekolah_tanpa_salesman'].append(name)

used = set()
for nm, p in items.items():
    code = re.sub(r'[^A-Z0-9]+', '-', nm).strip('-')[:40]
    base, n = code, 1
    while code in used:
        n += 1
        code = f'{base}-{n}'
    used.add(code)
    if not p['sell']:
        FLAGS['item_tanpa_harga_jual'].append(nm)
    sql.append(f"insert into md_items(code,name,type,category_id,base_unit_id,sale_price,purchase_price,legacy_code,updated_at) "
               f"select {sq(code)},{sq(nm)},'INVENTORY'::\"ErpItemType\",(select id from md_item_categories where code='{category(nm)}'),"
               f"(select id from md_units where code='{unit(nm)}'),{p['sell'] or 0},{p['buy'] or 0},'{TAG}',now() "
               f"where not exists (select 1 from md_items where code={sq(code)});")

# sinyal data mencurigakan
for name, c in custs.items():
    for _, ph in c['contacts']:
        if ph and re.fullmatch(r'0831\d{4}077\d|0831\d{3}0?\d*', ph or '') and ph.startswith('08311310'):
            FLAGS['telepon_urut_mencurigakan'].append(f'{name}: {ph}')
byphone = collections.defaultdict(set)
for name, c in custs.items():
    for _, ph in c['contacts']:
        if ph:
            byphone[ph].add(name)
for ph, ns in byphone.items():
    if len(ns) > 1:
        FLAGS['telepon_sama_beda_sekolah'].append(f'{ph}: {sorted(ns)}')

sql.append('commit;' if APPLY else 'rollback;')
out = psql('\n'.join(sql))
verify = psql(
    f"select 'partners', count(*) from md_partners where legacy_code='{TAG}' union all "
    f"select 'customers', count(*) from md_partners where legacy_code='{TAG}' and code like 'CUST-%' union all "
    f"select 'salesman', count(*) from md_partners where legacy_code='{TAG}' and code like 'SLS-%' union all "
    f"select 'addresses', count(*) from md_partner_addresses a join md_partners p on p.id=a.partner_id where p.legacy_code='{TAG}' union all "
    f"select 'contacts', count(*) from md_partner_contacts a join md_partners p on p.id=a.partner_id where p.legacy_code='{TAG}' union all "
    f"select 'items', count(*) from md_items where legacy_code='{TAG}' union all "
    f"select 'categories', count(*) from md_item_categories where legacy_code='{TAG}'", True)
print('MODE:', 'APPLY (committed)' if APPLY else 'DRY-RUN (rolled back)')
print('rencana: customers', len(custs), '| items', len(items), '| salesman', len(SALESMEN))
print('DB setelah run:\n' + verify)
with open(os.path.join(HERE, 'import-flags.json'), 'w') as f:
    json.dump(FLAGS, f, indent=1, ensure_ascii=False)
for k, v in FLAGS.items():
    print(f'[{k}] {len(v)}')

ROLLBACK_SQL = """
delete from md_partner_contacts where partner_id in (select id from md_partners where legacy_code='data-client');
delete from md_partner_addresses where partner_id in (select id from md_partners where legacy_code='data-client');
update md_partners set salesman_id=null where legacy_code='data-client';
delete from md_partners where legacy_code='data-client';
delete from md_items where legacy_code='data-client';
delete from md_item_categories where legacy_code='data-client';
"""
